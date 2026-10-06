"""Accuracy tests for collection and counting - no network or tokens needed.

A fake GitHub client answers the API calls the collectors make, so every edge
case (pagination, failures, search limits, duplicates, periods, time zones)
is exercised deterministically.

Run directly:

    python tests/test_accuracy.py

or with pytest:

    pytest -q tests
"""

from __future__ import annotations

import asyncio
import re
import sys
import tempfile
from dataclasses import replace
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

# Make the package importable when run as a plain script.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from github_contrib.client import GitHubClient, GitHubError  # noqa: E402
from github_contrib.commits import (  # noqa: E402
    collect_commits_for_repo,
    collect_pr_commits,
    enrich_commits_with_stats,
)
from github_contrib.config import AppConfig, resolve_period  # noqa: E402
from github_contrib.discovery import search_all  # noqa: E402
from github_contrib.filters import (  # noqa: E402
    apply_date_range,
    dedupe_contributions,
    localize,
    prepare_report_data,
)
from github_contrib.models import (  # noqa: E402
    CollectedData,
    CommitRecord,
    Coverage,
    PullRequestRecord,
    RepoRecord,
)

UTC = timezone.utc


# ---------------------------------------------------------------------------
# Fake GitHub client
# ---------------------------------------------------------------------------

class FakeClient:
    """Answers ``request`` / ``paginate`` / ``get_json`` / ``search`` from a routing table.

    ``routes`` maps ``(path, frozenset(params))`` or just ``path`` to either a
    list of items (paginate), a payload (request), or an exception to raise.
    """

    _parse_next_link = staticmethod(GitHubClient._parse_next_link)

    def __init__(self, routes: dict, search=None) -> None:
        self.routes = routes
        self.search_fn = search
        self.calls: list[tuple[str, dict | None]] = []

    def _route(self, path: str, params: dict | None):
        self.calls.append((path, params))
        key = (path, frozenset((params or {}).items()))
        if key in self.routes:
            value = self.routes[key]
        elif path in self.routes:
            value = self.routes[path]
        else:
            return None
        if isinstance(value, Exception):
            raise value
        return value

    async def request(self, method, path, *, params=None, accept=None):
        value = self._route(path, params)
        if isinstance(value, tuple):  # (data, headers)
            return value[0], value[1], 200
        return value, {}, 200 if value is not None else 404

    async def get_json(self, path, *, params=None, accept=None):
        data, _headers, _status = await self.request("GET", path, params=params)
        return data

    async def paginate(self, path, *, params=None, accept=None, max_items=None):
        value = self._route(path, params)
        for item in value or []:
            yield item

    async def search(self, path, query, *, accept=None, split_above=None):
        self.calls.append((path, {"q": query}))
        return await self.search_fn(query, split_above)


def _payload(sha, *, login="neha", email="neha@x.com", date="2026-04-21T11:33:34Z",
             message="feat: thing", parents=1):
    return {
        "sha": sha,
        "html_url": f"https://github.com/acme/app/commit/{sha}",
        "commit": {
            "message": message,
            "author": {"name": "Neha", "email": email, "date": date},
            "committer": {"name": "GitHub", "email": "noreply@github.com", "date": date},
        },
        "author": {"login": login} if login else None,
        "parents": [{"sha": f"p{i}"} for i in range(parents)],
    }


def _repo(full_name="acme/app", *, default="main", fork=False, affiliated=True) -> RepoRecord:
    owner, name = full_name.split("/")
    return RepoRecord(full_name=full_name, name=name, owner=owner, organization=owner,
                      default_branch=default, is_fork=fork, affiliated=affiliated)


def _commit(sha, repo="acme/app", *, branch="main", when=datetime(2026, 4, 21, 11, 33, tzinfo=UTC),
            message="feat: thing", email="neha@x.com", parents=1, committed=None, adds=0) -> CommitRecord:
    owner, name = repo.split("/")
    return CommitRecord(
        repository=name, full_name=repo, owner=owner, organization=owner, sha=sha,
        message=message, message_first_line=message.splitlines()[0], author_login="neha",
        author_name="Neha", author_email=email, committer_name="Neha", committer_email=email,
        authored_date=when, committed_date=committed or when, branch=branch,
        url=f"https://github.com/{repo}/commit/{sha}", parent_count=parents,
        additions=adds, stats_fetched=bool(adds),
    )


def _pr(number, repo="acme/app", *, merged=True, merge_sha="", shas=(), created=None) -> PullRequestRecord:
    owner, name = repo.split("/")
    when = created or datetime(2026, 4, 20, tzinfo=UTC)
    return PullRequestRecord(
        repository=name, full_name=repo, organization=owner, number=number, title=f"PR {number}",
        author_login="neha", state="closed" if merged else "open", merged=merged,
        created_at=when, updated_at=when, closed_at=None, merged_at=when if merged else None,
        base_branch="main", head_branch=f"feature-{number}", url="", merge_commit_sha=merge_sha,
        commit_shas=list(shas),
    )


# ---------------------------------------------------------------------------
# Collection
# ---------------------------------------------------------------------------

def test_line_stats_are_read_merges_skipped_and_failures_reported():
    page2 = "https://api.github.com/repos/acme/app/commits/c1?page=2"
    routes = {
        "/repos/acme/app/commits/c1": (
            {"stats": {"additions": 10, "deletions": 4}, "files": [{}] * 300},
            {"Link": f'<{page2}>; rel="next"'},
        ),
        page2: ({"files": [{}] * 20}, {}),
        "/repos/acme/app/commits/c2": GitHubError("boom"),
    }
    client = FakeClient(routes)
    c1, c1_fork = _commit("c1"), _commit("c1", repo="neha/app-fork")
    c2, merge = _commit("c2"), _commit("m1", parents=2)
    coverage = Coverage()
    asyncio.run(enrich_commits_with_stats({"acme/app": client}, [c1, c1_fork, c2, merge], coverage))

    assert (c1.additions, c1.deletions, c1.files_changed, c1.stats_fetched) == (10, 4, 320, True)
    assert (c1_fork.additions, c1_fork.stats_fetched) == (10, True), "same SHA shares its stats"
    assert c2.stats_fetched is False and c2.additions == 0
    assert merge.stats_fetched is False
    stat_calls = [path for path, _ in client.calls if "/commits/" in path]
    assert stat_calls.count("/repos/acme/app/commits/c1") == 1, "one request per SHA"
    assert not any(path.endswith("/m1") for path in stat_calls), "merge commits are not fetched"
    assert any("1 commit(s)" in note and "acme/app@c2" in note for note in coverage.notes)
    print("ok  test_line_stats_are_read_merges_skipped_and_failures_reported")


def test_branch_scanning_dedupes_heads_and_reports_failures():
    branches = [
        {"name": "zeta", "commit": {"sha": "h-main"}},  # same head as main: not rescanned
        {"name": "main", "commit": {"sha": "h-main"}},
        {"name": "feature", "commit": {"sha": "h-feat"}},
        {"name": "broken", "commit": {"sha": "h-broken"}},
    ]

    def commits_for(branch):
        return ("/repos/acme/app/commits", frozenset({"author": "neha", "sha": branch}.items()))

    routes = {
        "/repos/acme/app/branches": branches,
        commits_for("main"): [_payload("a"), _payload("b")],
        commits_for("feature"): [_payload("b"), _payload("c")],
        commits_for("broken"): GitHubError("Network error"),
    }
    client = FakeClient(routes)
    coverage = Coverage()
    commits = asyncio.run(collect_commits_for_repo(
        client, _repo(), ["neha"], scan_all_branches=True, coverage=coverage
    ))
    assert [(c.sha, c.branch) for c in commits] == [("a", "main"), ("b", "main"), ("c", "feature")]
    scanned = {dict(params)["sha"] for path, params in client.calls if path.endswith("/commits")}
    assert scanned == {"main", "feature", "broken"}, scanned  # zeta shares main's head
    assert any("branch 'broken'" in note for note in coverage.notes), coverage.notes

    # Upstream projects (not affiliated) only have their default branch scanned.
    client2 = FakeClient(routes)
    commits2 = asyncio.run(collect_commits_for_repo(
        client2, _repo(affiliated=False), ["neha"], scan_all_branches=True, coverage=Coverage()
    ))
    assert [c.sha for c in commits2] == ["a", "b"]
    assert all(path != "/repos/acme/app/branches" for path, _ in client2.calls)
    print("ok  test_branch_scanning_dedupes_heads_and_reports_failures")


def test_search_splits_ranges_beyond_the_1000_cap():
    # 2500 commits spread over 2020-2024; a query returns at most 1000.
    start = datetime(2020, 1, 1, tzinfo=UTC)
    stamps = [start + timedelta(hours=17 * i) for i in range(2500)]

    async def search(query, split_above):
        lo, hi = re.search(r"author-date:(\S+)\.\.(\S+)", query).groups()
        lo_dt, hi_dt = datetime.fromisoformat(lo), datetime.fromisoformat(hi)
        hits = [{"sha": s.isoformat(), "repository": {"full_name": "acme/app"}}
                for s in stamps if lo_dt <= s <= hi_dt]
        if split_above is not None and len(hits) > split_above:
            return hits[:100], len(hits), False
        return hits[:1000], len(hits), False

    client = FakeClient({}, search=search)
    coverage = Coverage()
    items = asyncio.run(search_all(
        client, "/search/commits", "author:neha", "author-date",
        datetime(2019, 1, 1, tzinfo=UTC), datetime(2026, 1, 1, tzinfo=UTC), coverage, "commits by neha",
    ))
    assert len(items) == 2500 and len({i["sha"] for i in items}) == 2500
    assert coverage.notes == []

    async def flaky(query, split_above):
        return [{"sha": "x"}], 1, True

    coverage2 = Coverage()
    asyncio.run(search_all(FakeClient({}, search=flaky), "/search/commits", "author:neha", "author-date",
                           start, start + timedelta(days=1), coverage2, "commits by neha"))
    assert any("timed out" in note for note in coverage2.notes)
    print("ok  test_search_splits_ranges_beyond_the_1000_cap")


def test_pr_commits_recover_deleted_branch_work():
    routes = {
        "/repos/acme/app/pulls/7/commits": [
            _payload("p1"),
            _payload("p2", login=None, email="neha@work.com"),  # unlinked email
            _payload("p3", login="someone-else", email="other@x.com"),
        ],
    }
    coverage = Coverage()
    shas, records = asyncio.run(collect_pr_commits(
        FakeClient(routes), _repo(), _pr(7), ["neha"], ["neha@work.com"], coverage
    ))
    assert shas == ["p1", "p2", "p3"]
    assert [(c.sha, c.branch, c.author_login) for c in records] == [
        ("p1", "feature-7", "neha"), ("p2", "feature-7", "neha"),
    ]
    routes_big = {"/repos/acme/app/pulls/8/commits": [_payload(f"s{i}") for i in range(250)]}
    asyncio.run(collect_pr_commits(FakeClient(routes_big), _repo(), _pr(8), ["neha"], [], coverage))
    assert any("first 250 commits" in note for note in coverage.notes)
    print("ok  test_pr_commits_recover_deleted_branch_work")


# ---------------------------------------------------------------------------
# Counting
# ---------------------------------------------------------------------------

def test_dedupe_counts_each_change_once():
    repos = [_repo("acme/app"), _repo("neha/app", fork=True)]
    t = datetime(2026, 5, 1, 9, 0, tzinfo=UTC)
    commits = [
        # Same SHA in a fork and its upstream: the upstream copy is kept.
        _commit("s1", repo="neha/app", when=t),
        _commit("s1", repo="acme/app", when=t),
        # Rebased copy: same author/time/subject, new SHA; the default-branch copy stays.
        _commit("r-orig", branch="topic", when=t + timedelta(hours=1), message="fix: login",
                committed=t + timedelta(hours=1)),
        _commit("r-main", branch="main", when=t + timedelta(hours=1), message="fix: login",
                committed=t + timedelta(days=2)),
        # Squash merge: originals o1, o2 folded into the squash commit sq.
        _commit("o1", branch="feature-9", when=t + timedelta(hours=2), message="wip 1"),
        _commit("o2", branch="feature-9", when=t + timedelta(hours=3), message="wip 2"),
        _commit("sq", branch="main", when=t + timedelta(days=1), message="Add search (#9)"),
        # Merge-commit merge: originals stay in history and keep counting.
        _commit("k1", branch="main", when=t + timedelta(hours=4), message="keep 1"),
        _commit("mc", branch="main", when=t + timedelta(days=1, hours=1), message="Merge #10", parents=2),
        # Same subject and author but another time: a different change.
        _commit("d1", branch="main", when=t + timedelta(hours=5), message="fix: login"),
    ]
    prs = [
        _pr(9, merge_sha="sq", shas=["o1", "o2"]),
        _pr(10, merge_sha="mc", shas=["k1"]),
    ]
    data = CollectedData(repos=repos, commits=commits, pull_requests=prs)
    out, result = dedupe_contributions(data)
    kept = sorted((c.full_name, c.sha) for c in out.commits)
    assert kept == sorted([
        ("acme/app", "s1"), ("acme/app", "r-main"), ("acme/app", "sq"),
        ("acme/app", "k1"), ("acme/app", "mc"), ("acme/app", "d1"),
    ]), kept
    assert (result.same_commit, result.rebased_copies, result.squashed) == (1, 1, 2)
    # Idempotent: running again changes nothing.
    again, result2 = dedupe_contributions(out)
    assert len(again.commits) == len(out.commits) and result2.total == 0
    print("ok  test_dedupe_counts_each_change_once")


def test_exclusions_apply_before_dedupe():
    """A commit pushed to both a personal repo and the company repo must stay
    counted (in the company repo) when personal repos are excluded."""
    repos = [_repo("neha/side"), _repo("acme/app")]
    commits = [_commit("same", repo="neha/side"), _commit("same", repo="acme/app")]
    config = AppConfig(accounts=[], target_logins=["neha"], exclude_own_repos=True)
    out, scope = prepare_report_data(CollectedData(repos=repos, commits=commits), config)
    assert [(c.full_name, c.sha) for c in out.commits] == [("acme/app", "same")]
    assert scope.excluded["commits"] == 1 and scope.duplicates.total == 0
    print("ok  test_exclusions_apply_before_dedupe")


def test_period_boundaries_are_inclusive_in_the_report_time_zone():
    since, until, tz, name = resolve_period("2026-07-01", "2026-07-31", "Asia/Kolkata")
    ist = ZoneInfo("Asia/Kolkata")
    inside_start = datetime(2026, 7, 1, 0, 0, tzinfo=ist)  # = 2026-06-30 18:30 UTC
    inside_end = datetime(2026, 7, 31, 23, 59, 59, tzinfo=ist)
    before = inside_start - timedelta(seconds=1)
    after = datetime(2026, 8, 1, 0, 0, tzinfo=ist)
    commits = [_commit(s, when=w) for s, w in
               [("start", inside_start), ("end", inside_end), ("before", before), ("after", after)]]
    commits.append(replace(_commit("undated"), authored_date=None))
    prs = [_pr(1, created=inside_end), _pr(2, created=after)]
    data = CollectedData(repos=[_repo()], commits=commits, pull_requests=prs)
    out, removed = apply_date_range(data, since, until)
    assert sorted(c.sha for c in out.commits) == ["end", "start"]
    assert [p.number for p in out.pull_requests] == [1]
    assert removed == {"commits": 3, "pull_requests": 1}
    assert name == "Asia/Kolkata" and str(tz) == "Asia/Kolkata"
    print("ok  test_period_boundaries_are_inclusive_in_the_report_time_zone")


def test_days_follow_the_report_time_zone():
    from github_contrib.insights import compute_insights
    from github_contrib.statistics import compute_statistics

    # Monday 2026-04-20 20:30 UTC is Tuesday 02:00 in India.
    data = CollectedData(repos=[_repo()], commits=[_commit("a", when=datetime(2026, 4, 20, 20, 30, tzinfo=UTC))])
    ist = localize(data, ZoneInfo("Asia/Kolkata"))
    stats = compute_statistics(ist)
    tuesday = stats.commits_by_weekday.set_index("weekday").loc["Tuesday", "commits"]
    assert int(tuesday) == 1
    assert compute_insights(ist).busiest_day == "Tuesday"
    assert str(stats.summary_dict["first_contribution_date"]).startswith("2026-04-21T02:00:00+05:30")
    print("ok  test_days_follow_the_report_time_zone")


def test_line_totals_exclude_merges_and_average_over_measured_commits():
    from github_contrib.statistics import compute_statistics

    commits = [_commit("a", adds=100), _commit("b", adds=50), _commit("m", parents=2)]
    stats = compute_statistics(CollectedData(repos=[_repo()], commits=commits))
    assert stats.summary_dict["total_lines_added"] == 150
    assert stats.summary_dict["avg_lines_added_per_commit"] == 75.0  # not 150 / 3
    print("ok  test_line_totals_exclude_merges_and_average_over_measured_commits")


# ---------------------------------------------------------------------------
# Exports and report safety
# ---------------------------------------------------------------------------

def test_new_fields_roundtrip_and_formulas_are_neutralized():
    from github_contrib.exporters import export_csvs, export_excel
    from github_contrib.offline import load_collected_from_csv
    from github_contrib.statistics import compute_statistics

    evil = '=HYPERLINK("https://evil.example","x")'
    commits = [_commit("a", message=evil, parents=2), _commit("b", adds=7)]
    prs = [_pr(3, merge_sha="b", shas=["a", "b"])]
    repo = replace(_repo(affiliated=False), description="+cmd|' /C calc'!A0")
    data = CollectedData(repos=[repo], commits=commits, pull_requests=prs)
    stats = compute_statistics(data)
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp)
        export_csvs(out, stats)
        raw = (out / "commits.csv").read_text(encoding="utf-8-sig")
        assert "'=HYPERLINK" in raw, "CSV text must not start with ="
        loaded = load_collected_from_csv(out)
        by_sha = {c.sha: c for c in loaded.commits}
        assert by_sha["a"].message == evil and by_sha["a"].parent_count == 2
        assert by_sha["b"].stats_fetched and by_sha["b"].additions == 7
        assert loaded.pull_requests[0].merge_commit_sha == "b"
        assert loaded.pull_requests[0].commit_shas == ["a", "b"]
        assert loaded.repos[0].affiliated is False and loaded.repos[0].description.startswith("+cmd")

        xlsx = export_excel(out, stats)
        from openpyxl import load_workbook

        sheet = load_workbook(xlsx)["Commits"]
        cells = [cell for row in sheet.iter_rows() for cell in row if cell.value == evil]
        assert cells and all(cell.data_type == "s" for cell in cells), "Excel must store text, not a formula"
    print("ok  test_new_fields_roundtrip_and_formulas_are_neutralized")


def test_report_html_is_locked_down():
    import base64
    import hashlib

    from github_contrib.filters import ReportScope
    from github_contrib.htmlreport import _SCRIPT, render_html, render_markdown
    from github_contrib.insights import compute_insights
    from github_contrib.statistics import compute_statistics

    bad = replace(_commit("x", message="<script>alert(1)</script>"), url="javascript:alert(1)")
    data = CollectedData(repos=[_repo()], commits=[bad, _commit("y")])
    stats = compute_statistics(data)
    insights = compute_insights(data)
    since, until, _tz, name = resolve_period("2026-01-01", "2026-12-31", "Asia/Kolkata")
    scope = ReportScope(since=since, until=until, timezone=name, notes=["acme/app: branch 'x' failed"])
    html = render_html(stats, insights, stats.summary_dict, Path("nowhere"), include_charts=False, scope=scope)

    digest = base64.b64encode(hashlib.sha256(_SCRIPT.encode()).digest()).decode()
    assert f"script-src 'sha256-{digest}'" in html and "default-src 'none'" in html
    assert "<script>alert(1)</script>" not in html and "&lt;script&gt;" in html
    assert "javascript:" not in html
    assert "1 Jan 2026 – 31 Dec 2026 (Asia/Kolkata)" in html and "Commits in period" in html
    assert "Data completeness" in html and "branch &#x27;x&#x27; failed" in html
    assert "How this report was compiled" in html
    md = render_markdown(stats, insights, stats.summary_dict, scope)
    assert "<script>" not in md and "**Period:** 1 Jan 2026" in md
    print("ok  test_report_html_is_locked_down")


def _all_tests():
    return [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]


def main() -> int:
    failures = 0
    for test in _all_tests():
        try:
            test()
        except Exception as exc:  # noqa: BLE001
            failures += 1
            print(f"FAIL {test.__name__}: {exc!r}")
            import traceback
            traceback.print_exc()
    total = len(_all_tests())
    print(f"\n{total - failures}/{total} tests passed.")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
