"""Insights & analytics tests - synthetic data, no network.

Run directly:

    python tests/test_analytics.py

or with pytest:

    pytest -q tests
"""

from __future__ import annotations

import sys
import tempfile
from datetime import datetime, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

# Make the package importable when run as a plain script.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from github_contrib.analytics import (  # noqa: E402
    WORK_TYPES,
    classify_commit,
    clean_title,
    compute_analytics,
    human_duration,
    percentile,
)
from github_contrib.exporters import export_csvs, export_excel  # noqa: E402
from github_contrib.graphql import pr_record  # noqa: E402
from github_contrib.htmlreport import render_html, render_markdown  # noqa: E402
from github_contrib.insights import compute_insights  # noqa: E402
from github_contrib.models import CollectedData, CommitRecord, PullRequestRecord, RepoRecord  # noqa: E402
from github_contrib.offline import load_collected_from_csv  # noqa: E402
from github_contrib.pull_requests import _pr_from_payload  # noqa: E402
from github_contrib.statistics import attach_analytics, compute_statistics  # noqa: E402

UTC = timezone.utc
IST = ZoneInfo("Asia/Kolkata")


def _commit(sha: str, when: datetime, message: str = "feat: thing", *, repo: str = "acme/app",
            login: str = "neha", adds: int = 0, dels: int = 0, parents: int = 1) -> CommitRecord:
    owner, name = repo.split("/")
    return CommitRecord(
        repository=name, full_name=repo, owner=owner, organization=owner if owner != login else "",
        sha=sha, message=message, message_first_line=message.splitlines()[0], author_login=login,
        author_name="Neha", author_email="neha@x.com", committer_name="Neha", committer_email="neha@x.com",
        authored_date=when, committed_date=when, branch="main", url=f"https://github.com/{repo}/commit/{sha}",
        additions=adds, deletions=dels, stats_fetched=bool(adds or dels), parent_count=parents,
    )


def _pr(number: int, created: datetime, merged: datetime | None = None, *, state: str = "closed",
        title: str = "Add login screen", repo: str = "acme/app", adds: int | None = None,
        dels: int | None = None, files: int | None = None, shas: tuple = ()) -> PullRequestRecord:
    owner, name = repo.split("/")
    return PullRequestRecord(
        repository=name, full_name=repo, organization=owner, number=number, title=title,
        author_login="neha", state=state, merged=merged is not None, created_at=created,
        updated_at=merged or created, closed_at=merged, merged_at=merged, base_branch="main",
        head_branch=f"feature/{number}", url=f"https://github.com/{repo}/pull/{number}",
        commit_shas=list(shas), additions=adds, deletions=dels, changed_files=files,
    )


def _data(commits, prs=(), logins=("neha",)) -> CollectedData:
    repos = sorted({c.full_name for c in commits} | {p.full_name for p in prs})
    return CollectedData(
        repos=[RepoRecord(full_name=r, name=r.split("/")[1], owner=r.split("/")[0], organization="") for r in repos],
        commits=list(commits), pull_requests=list(prs), meta={"target_logins": list(logins)},
    )


def test_commit_messages_are_sorted_into_kinds_of_work():
    cases = {
        "feat(auth): add OTP login": "New features",
        "fix: crash on empty cart": "Bug fixes",
        "refactor: split payment service": "Refactoring & cleanup",
        "docs: explain the deploy steps": "Documentation",
        "chore(deps): bump axios": "Build, CI & chores",
        "test: cover the parser": "Tests",
        "perf: cache the dashboard query": "Improvements",
        "[ABC-12] Fix toast alignment": "Bug fixes",
        "ABC-12: Add dark mode": "New features",
        "Razorpay amount bug fixed": "Bug fixes",
        "Trading UI issues resolved": "Bug fixes",
        "Staging Deployment test commit": "Build, CI & chores",  # a deploy, not a test
        "Add unit tests for the parser": "Tests",
        "Update README.md": "Documentation",
        "Updated theme": "Improvements",
        "My request screen added": "New features",
        "Bookmark functionality for stock screens": "New features",
        "tenant if function moved to controller": "Refactoring & cleanup",
        "Revert \"Add banner\"": "Build, CI & chores",
        "wip: tidy the imports": "Refactoring & cleanup",  # unknown prefix: judged by the rest
        "app(login): add button": "New features",
        # Changelog-style prefixes are kinds of work too.
        "Added: login page": "New features",
        "Updated: header colours": "Improvements",
        "Fixed: crash on launch": "Bug fixes",
        "New: dark mode": "New features",
        # Housekeeping, not fixes (they mention "merge", "issue").
        "Resolved merge conflicts": "Build, CI & chores",
        "Add issue templates": "Build, CI & chores",
        "Help and support complete": "Other",
        "": "Other",
    }
    for message, kind in cases.items():
        assert classify_commit(message) == kind, (message, classify_commit(message), kind)
    for merge in ("Merge pull request #1 from neha/dev", "Merge branch 'main' into feature", "Merged PR 12: x",
                  "Merged in feature/x (pull request #12)"):
        assert classify_commit(merge) is None, merge
    assert set(cases.values()) <= set(WORK_TYPES)
    print("ok  test_commit_messages_are_sorted_into_kinds_of_work")


def test_helpers():
    assert percentile([], 50) is None
    assert percentile([5.0], 90) == 5.0
    assert percentile([1.0, 2.0, 3.0, 4.0], 50) == 2.5
    assert percentile([0.0, 10.0], 90) == 9.0
    assert human_duration(None) == "n/a" and human_duration(0.001) == "under a minute"
    assert human_duration(0.75) == "45 min" and human_duration(5.5) == "5.5 h" and human_duration(24) == "24 h"
    assert human_duration(72) == "3 days" and human_duration(60) == "2.5 days"
    assert clean_title("feat(ui): add dark mode") == "Add dark mode"
    assert clean_title("[ABC-1]  fix   the  toast") == "Fix the toast"
    print("ok  test_helpers")


def test_delivery_measures_time_to_merge_and_size():
    t0 = datetime(2026, 3, 2, 10, 0, tzinfo=UTC)
    prs = [
        _pr(1, t0, t0 + timedelta(minutes=30), adds=4, dels=2, files=1),       # XS, within 1 h
        _pr(2, t0, t0 + timedelta(hours=5), adds=60, dels=20, files=3),        # S, 1-24 h
        _pr(3, t0, t0 + timedelta(days=2), adds=300, dels=50, files=9),        # M, 1-3 days
        _pr(4, t0, t0 + timedelta(days=10), adds=1500, dels=200, files=40, title="feat: payments v2"),  # XL
        _pr(5, t0, None, state="closed"),                                      # closed unmerged
        _pr(6, t0, None, state="open"),                                        # still open
        _pr(7, t0, t0 + timedelta(hours=1)),                                   # merged, size unknown
    ]
    d = compute_analytics(_data([_commit("a", t0)], prs)).delivery
    assert (d.opened, d.merged, d.closed_unmerged, d.still_open) == (7, 5, 1, 1)
    assert abs(d.merge_rate - 5 / 6) < 1e-9
    assert d.hours_to_merge == sorted([0.5, 5, 48, 240, 1])
    assert d.median_hours == 5 and d.within_day == 3 and d.within_week == 4
    assert dict(d.time_buckets) == {"Within 1 hour": 2, "1–24 hours": 1, "1–3 days": 1, "3–7 days": 0,
                                    "1–4 weeks": 1, "Over 4 weeks": 0}
    assert d.sized == 4 and d.median_size == (80 + 350) / 2
    assert dict(d.size_buckets) == {"XS · under 10 lines": 1, "S · 10–99": 1, "M · 100–499": 1,
                                    "L · 500–999": 0, "XL · 1,000+": 1}
    assert [p.number for p in d.biggest] == [4, 3, 2, 1]
    assert compute_analytics(_data([_commit("a", t0)])).delivery is None, "no pull requests: no delivery"
    print("ok  test_delivery_measures_time_to_merge_and_size")


def test_rhythm_consistency_and_best_30_days():
    # Wednesday 14:30 IST x3, Saturday 23:10 IST x1, then a burst in May.
    base = datetime(2026, 1, 7, 14, 30, tzinfo=IST)  # a Wednesday
    commits = [_commit(f"w{i}", base + timedelta(weeks=i)) for i in range(3)]
    commits.append(_commit("sat", datetime(2026, 1, 10, 23, 10, tzinfo=IST)))
    burst = datetime(2026, 5, 4, 11, 0, tzinfo=IST)  # a Monday
    commits += [_commit(f"b{i}", burst + timedelta(days=i)) for i in range(10)]
    r = compute_analytics(_data(commits)).rhythm
    assert r.total == 14
    assert r.heatmap[2][14] == 3 and r.heatmap[5][23] == 1, "weekday x hour in the data's own time zone"
    assert abs(r.weekend_share - 3 / 14) < 1e-9, "the Saturday plus the burst's Sat and Sun"
    assert r.best_30_commits == 10 and r.best_30_start.isoformat() == "2026-05-04"
    assert r.best_30_end.isoformat() == "2026-05-13"
    assert r.weeks_in_span == 19 and r.active_weeks == 5
    assert r.peak_start in range(24) and 0 < r.peak_share <= 1
    assert compute_analytics(_data([])).rhythm is None
    print("ok  test_rhythm_consistency_and_best_30_days")


def test_quarters_compare_only_complete_ones_and_list_what_shipped():
    def at(y, m, d):
        return datetime(y, m, d, 12, tzinfo=UTC)

    commits = [_commit(f"a{i}", at(2026, 2, 10)) for i in range(2)]                       # Q1 (work starts here)
    commits += [_commit(f"b{i}", at(2026, 5, 3), "Updated the layout", adds=10, dels=5) for i in range(4)]  # Q2
    commits += [_commit(f"c{i}", at(2026, 8, 3), "fix: toast") for i in range(6)]             # Q3
    commits += [_commit("d0", at(2026, 10, 1), "Add ferry ladder logic for class placement")]  # Q4 (in progress)
    prs = [
        _pr(1, at(2026, 4, 1), at(2026, 4, 2), title="Add dark mode", adds=500, dels=10),
        _pr(2, at(2026, 4, 3), at(2026, 4, 4), title="feat: OTP login flow", adds=40, dels=10),
        _pr(3, at(2026, 4, 5), at(2026, 4, 6), title="Commit"),        # one word: says nothing
        _pr(4, at(2026, 9, 28), at(2026, 10, 20), title="Late merge of the payments page"),  # merged after the period
    ]
    a = compute_analytics(_data(commits, prs), until=datetime(2026, 10, 15, tzinfo=UTC))
    labels = [q.label for q in a.quarters]
    assert labels == ["2026 Q1", "2026 Q2", "2026 Q3", "2026 Q4"], labels
    q1, q2, q3, q4 = a.quarters
    assert q1.partial and q4.partial and not q2.partial and not q3.partial
    assert (q2.commits, q2.prs_opened, q2.prs_merged, q2.repositories, q2.active_days) == (4, 3, 3, 1, 1)
    assert (q2.additions, q2.deletions) == (40, 20)
    assert q2.change is None, "Q1 is partial (the work began mid-quarter): no comparison"
    assert q3.change == 6 / 4 - 1 and q4.change is None
    assert q3.prs_opened == 1 and q4.prs_merged == 0, "a merge after the period belongs to no quarter"
    assert [what for what, _repo, _ref in q2.shipped] == ["Add dark mode", "OTP login flow"]
    assert q2.shipped[0] == ("Add dark mode", "app", "PR #1")
    # No merged pull requests: the commits that describe features, improvements or fixes.
    assert q4.shipped == [("Add ferry ladder logic for class placement", "app", "")]
    # A quiet quarter in between is shown, not skipped.
    gap = compute_analytics(_data([_commit("x", at(2026, 1, 5)), _commit("y", at(2026, 7, 5))]))
    assert [q.label for q in gap.quarters] == ["2026 Q1", "2026 Q2", "2026 Q3"] and gap.quarters[1].commits == 0
    print("ok  test_quarters_compare_only_complete_ones_and_list_what_shipped")


def test_a_running_period_counts_weeks_and_quarters_up_to_today():
    def at(y, m, d):
        return datetime(y, m, d, 12, tzinfo=UTC)

    # Every other week from Mon 5 Jan, and one commit in October.
    commits = [_commit(f"w{i}", at(2026, 1, 5) + timedelta(weeks=i)) for i in range(0, 40, 2)]
    commits.append(_commit("q4", at(2026, 10, 7)))
    year = {"since": datetime(2026, 1, 1, tzinfo=UTC), "until": datetime(2026, 12, 31, 23, 59, tzinfo=UTC)}
    a = compute_analytics(_data(commits), **year, now=at(2026, 10, 15))
    q4 = a.quarters[-1]
    assert q4.label == "2026 Q4" and q4.partial and q4.change is None, "15 days of Q4 aren't compared with Q3"
    r = a.rhythm
    # The year so far: Thu 1 Jan to Thu 15 Oct, 42 weeks (Mon 29 Dec to Mon 12 Oct).
    assert (r.span_start.isoformat(), r.span_end.isoformat()) == ("2026-01-01", "2026-10-15")
    assert (r.weeks_in_span, r.active_weeks) == (42, 21)
    assert "Committed code in 21 of 42 weeks (50%)." in a.highlights
    # "From 1 Jan" runs to today too; all time runs from the first to the latest commit.
    since_only = compute_analytics(_data(commits), since=year["since"], now=at(2026, 10, 15)).rhythm
    assert since_only.span_end.isoformat() == "2026-10-15"
    all_time = compute_analytics(_data(commits), now=at(2026, 10, 15)).rhythm
    assert (all_time.span_start.isoformat(), all_time.span_end.isoformat()) == ("2026-01-05", "2026-10-07")
    # Under half the weeks: a fact in the section, not a highlight.
    sparse = compute_analytics(_data(commits[::3]), **year, now=at(2026, 10, 15))
    assert not any(h.startswith("Committed code in") for h in sparse.highlights), sparse.highlights
    # A merged pull request with no opening time and no commits: no crash.
    lone = compute_analytics(_data([], [_pr(9, None, at(2026, 2, 1))]))
    assert [(q.label, q.prs_merged) for q in lone.quarters] == [("2026 Q1", 1)]
    print("ok  test_a_running_period_counts_weeks_and_quarters_up_to_today")


def test_timeline_shows_when_each_repository_was_worked_on():
    def at(y, m, d):
        return datetime(y, m, d, 12, tzinfo=UTC)

    commits = [_commit("a1", at(2026, 3, 2), repo="acme/api"), _commit("a2", at(2026, 3, 20), repo="acme/api"),
               _commit("a3", at(2026, 6, 1), repo="acme/api")]
    commits += [_commit(f"w{i}", at(2026, 1, 10 + i), repo="acme/web") for i in range(5)]
    commits += [_commit(f"s{i}", at(2026, 2, 1), repo=f"neha/small{i}") for i in range(12)]
    data = _data(commits)
    a = compute_analytics(data, since=at(2026, 1, 1), until=at(2026, 12, 31), now=at(2026, 9, 30))
    t = a.timeline
    assert len(t.rows) == 12 and t.hidden == 2, "the 12 repositories with the most commits"
    names = [row.repository for row in t.rows]
    assert names[0] == "acme/web" and names[-1] == "acme/api", "in the order the work began"
    api = t.rows[-1]
    assert (api.first.isoformat(), api.last.isoformat(), api.commits) == ("2026-03-02", "2026-06-01", 3)
    assert api.by_month == {"2026-03": 2, "2026-06": 1}
    assert (t.start.isoformat(), t.end.isoformat()) == ("2026-01-10", "2026-09-30"), "up to today, in the period"
    stats = compute_statistics(data)
    attach_analytics(stats, a)
    insights = compute_insights(data)
    with tempfile.TemporaryDirectory() as tmp:
        page = render_html(stats, insights, stats.summary_dict, Path(tmp), include_charts=False)
    assert "<h3>Project timeline</h3>" in page and "2 more repositories with fewer commits are left off." in page
    assert "title='Mar 2026 · 2 commits'" in page and ">Apr 2026</div>" in page and ">Jul 2026</div>" in page
    md = render_markdown(stats, insights, stats.summary_dict)
    assert "### Project timeline" in md and "months with commits" in md
    # One repository: nothing to compare, no timeline.
    solo = _data(commits[:3])
    stats = compute_statistics(solo)
    attach_analytics(stats, compute_analytics(solo))
    with tempfile.TemporaryDirectory() as tmp:
        assert "Project timeline" not in render_html(stats, compute_insights(solo), stats.summary_dict, Path(tmp),
                                                     include_charts=False)
    print("ok  test_timeline_shows_when_each_repository_was_worked_on")


def test_footprint_and_work_mix_leave_merge_commits_out():
    t = datetime(2026, 3, 2, 10, tzinfo=UTC)
    commits = [_commit(f"s{i}", t, "fix: small", adds=3, dels=2) for i in range(8)]
    commits += [_commit("big", t, "feat: big", adds=900, dels=100)]
    commits += [_commit("m", t, "Merge pull request #3 from neha/dev", adds=5000, dels=5000, parents=2)]
    a = compute_analytics(_data(commits))
    f = a.footprint
    assert f.measured == 9 and f.median_lines == 5 and f.small_share == 8 / 9
    assert (f.additions, f.deletions) == (924, 116) and round(f.removed_per_100_added, 1) == 12.6
    assert a.work_mix.merges == 1 and a.work_mix.total == 9 and a.work_mix.counts["Bug fixes"] == 8
    assert compute_analytics(_data([_commit("n", t)])).footprint is None, "no line counts: no footprint"
    print("ok  test_footprint_and_work_mix_leave_merge_commits_out")


def test_highlights_and_summary_only_say_what_the_data_supports():
    t0 = datetime(2026, 1, 5, 10, tzinfo=UTC)
    commits = [
        _commit(f"c{i}", t0 + timedelta(days=3 * i), "feat: add screen" if i % 2 else "fix: bug",
                repo="acme/app" if i % 3 else "neha/side", adds=20, dels=5)
        for i in range(40)
    ]
    prs = [_pr(i, t0 + timedelta(days=i), t0 + timedelta(days=i, hours=2), adds=100, dels=20, files=4)
           for i in range(1, 6)]
    a = compute_analytics(_data(commits, prs), now=datetime(2026, 12, 31, tzinfo=UTC))
    text = "\n".join(a.highlights)
    assert "Merged 5 pull requests; half of them within 2 h of opening." in text
    assert "Most productive stretch:" in text and "Committed code in" in text
    assert "% of commits built or improved features; 50% fixed bugs." in text
    assert "Largest merged pull request:" in text and "+100 / −20 lines, 4 files" in text
    assert "went into repositories owned by organizations or other accounts" in text
    summary = a.summary()
    assert summary["median_hours_to_merge"] == 2.0 and summary["prs_merged_within_1_day"] == 5
    assert summary["median_pr_size_lines"] == 120 and summary["pr_merge_rate_pct"] == 100.0
    assert "New features 50%" in summary["work_mix"] and summary["median_commit_size_lines"] == 25
    # Little data, no pull requests: nothing that would overstate it.
    small = compute_analytics(_data(commits[:3]))
    assert small.delivery is None and not any("pull request" in h for h in small.highlights)
    assert not any("Most productive" in h for h in small.highlights), "fewer than 10 commits"
    print("ok  test_highlights_and_summary_only_say_what_the_data_supports")


def test_report_shows_the_section_and_workbook_has_the_sheets():
    t0 = datetime(2026, 1, 5, 10, tzinfo=UTC)
    commits = [_commit(f"c{i}", t0 + timedelta(days=5 * i), "feat: add screen", adds=10, dels=2) for i in range(30)]
    prs = [_pr(1, t0, t0 + timedelta(hours=3), title="Add the <b>screen</b>", adds=12, dels=2, files=2)]
    data = _data(commits, prs)
    stats = compute_statistics(data)
    attach_analytics(stats, compute_analytics(data))
    insights = compute_insights(data)
    with tempfile.TemporaryDirectory() as tmp:
        page = render_html(stats, insights, stats.summary_dict, Path(tmp), include_charts=False)
        assert "<section id='analytics'><h2>Insights &amp; analytics</h2>" in page
        assert "Highlights" in page and "Quarter by quarter" in page and "table class='heat'" in page
        assert "Add the &lt;b&gt;screen&lt;/b&gt;" in page and "<b>screen</b>" not in page, "titles are escaped"
        md = render_markdown(stats, insights, stats.summary_dict)
        assert "## Insights & analytics" in md and "### Work mix" in md
        workbook = export_excel(Path(tmp), stats)
        from openpyxl import load_workbook

        book = load_workbook(workbook, read_only=True)
        names = book.sheetnames
        book.close()  # Windows can't delete an open file
        assert {"Highlights", "Quarterly", "Work Mix", "PR Delivery", "Working Rhythm"} <= set(names), names
    # Without analytics (an older caller): no section, no error.
    plain = compute_statistics(data)
    with tempfile.TemporaryDirectory() as tmp:
        assert "id='analytics'" not in render_html(plain, insights, plain.summary_dict, Path(tmp), include_charts=False)
    print("ok  test_report_shows_the_section_and_workbook_has_the_sheets")


def test_pull_request_sizes_survive_every_path():
    repo = RepoRecord(full_name="acme/app", name="app", owner="acme", organization="acme")
    node = {"number": 7, "title": "Add x", "state": "MERGED", "createdAt": "2026-01-01T00:00:00Z",
            "mergedAt": "2026-01-02T00:00:00Z", "additions": 120, "deletions": 30, "changedFiles": 6}
    assert (pr_record(node, repo).additions, pr_record(node, repo).changed_files) == (120, 6)
    assert pr_record({**node, "additions": None}, repo).size is None
    rest = {"number": 8, "title": "y", "state": "closed", "merged_at": "2026-01-02T00:00:00Z",
            "created_at": "2026-01-01T00:00:00Z", "additions": 5, "deletions": 1, "changed_files": 2}
    assert _pr_from_payload(rest, repo).size == 6
    assert _pr_from_payload({k: v for k, v in rest.items() if k != "additions"}, repo).size is None
    # CSV and back (--regen): recorded sizes stay ints, unrecorded ones stay blank.
    t0 = datetime(2026, 1, 1, tzinfo=UTC)
    prs = [_pr(1, t0, t0 + timedelta(hours=1), adds=12, dels=3, files=2), _pr(2, t0, t0 + timedelta(hours=2))]
    stats = compute_statistics(_data([_commit("a", t0)], prs))
    with tempfile.TemporaryDirectory() as tmp:
        export_csvs(Path(tmp), stats)
        text = (Path(tmp) / "pull_requests.csv").read_text(encoding="utf-8-sig")
        assert ",12,3,2" in text and "12.0" not in text
        loaded = {p.number: p for p in load_collected_from_csv(Path(tmp)).pull_requests}
    assert (loaded[1].additions, loaded[1].deletions, loaded[1].changed_files) == (12, 3, 2)
    assert loaded[2].additions is None and loaded[2].size is None
    print("ok  test_pull_request_sizes_survive_every_path")


def _all_tests():
    return [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]


def main() -> int:
    failures = 0
    for test in _all_tests():
        try:
            test()
        except Exception as exc:  # noqa: BLE001
            failures += 1
            print(f"FAIL {test.__name__}: {exc!r}")
            import traceback
            traceback.print_exc()
    total = len(_all_tests())
    print(f"\n{total - failures}/{total} tests passed.")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
